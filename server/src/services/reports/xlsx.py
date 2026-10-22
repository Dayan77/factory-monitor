"""XLSX rendering with openpyxl (docs/reports.md §3.2).

* "Resumo" = the first page of the PDF; one sheet per data table; "Leia-me" with definitions and warnings.
* Raw data: fractions as numbers with format 0.0%; times in minutes (number) with the unit in the column name;
  timestamps in two columns (`<col>_local` Excel date/time in the plant time zone and `<col>_utc` ISO 8601).
* No formulas, no merged cells in data sheets. Undefined = empty cell + `motivo` column (never 0).
"""
from __future__ import annotations

import io
import re
from datetime import timezone
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import fmt
from .doc import NOTE_TEXT, Col, Doc, Table, coverage_note

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WARN_FILL = PatternFill("solid", fgColor="FFF4CE")
PCT_FMT = "0.0%"
SHEET_BAD = re.compile(r"[\[\]\*\?/\\:]")


def _sheet_name(wb: Workbook, name: str) -> str:
    base = SHEET_BAD.sub("_", name)[:31] or "Dados"
    n, k = base, 2
    while n in wb.sheetnames:
        suffix = f"_{k}"
        n = base[:31 - len(suffix)] + suffix
        k += 1
    return n


SNAKE = re.compile(r"^[a-z0-9_]+$")


def _header_cells(col: Col, human: bool = False) -> list[str]:
    """Data sheets: snake_case pt names (filters/pivots); 'Resumo': the PDF labels."""
    if human:
        name = col.label + (f" ({col.unit})" if col.unit else "")
        if col.kind == "datetime":
            return [f"{name} (local)", f"{name} (UTC)"]
        if col.kind == "min":
            return [f"{name} (min)"]
        return [name]
    name = col.xlsx_label or (col.label if SNAKE.match(col.label) else col.key)
    if col.kind == "datetime":
        return [f"{name}_local", f"{name}_utc"]
    if col.kind == "min" and not name.endswith("_min"):
        return [f"{name}_min"]
    if col.kind == "hours" and not name.endswith("_h"):
        return [f"{name}_h"]
    return [name]


def _cells(col: Col, value, tz: ZoneInfo) -> list:
    """Raw value(s) for one column; None stays empty."""
    if value is None:
        return [None, None] if col.kind == "datetime" else [None]
    if col.kind == "datetime":
        ts = fmt.parse_ts(value)
        return [ts.astimezone(tz).replace(tzinfo=None), ts.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")]
    if col.kind == "min":
        return [float(value) / 60.0]
    if col.kind in ("pct", "pp", "num0", "num1", "num2", "num6", "int", "hours"):
        if isinstance(value, bool):
            return [value]
        try:
            return [float(value) if col.kind != "int" else (int(value) if float(value).is_integer() else float(value))]
        except (TypeError, ValueError):
            return [value]
    if isinstance(value, bool):
        return ["sim" if value else "não"]
    return [value]


def write_table(ws, t: Table, tz: ZoneInfo, start_row: int = 1, title: bool = False, human: bool = False) -> int:
    r = start_row
    if title:
        ws.cell(row=r, column=1, value=t.title).font = Font(bold=True, size=12)
        r += 1
    headers: list[str] = []
    spans: list[tuple[Col, int]] = []
    for c in t.columns:
        hs = _header_cells(c, human)
        spans.append((c, len(hs)))
        headers += hs
    headers += ["motivo", "notas"]
    for j, h in enumerate(headers, start=1):
        cell = ws.cell(row=r, column=j, value=h)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    head_row = r
    r += 1
    for row in t.rows:
        j = 1
        reasons = row.get("_reason") or {}
        for c, n in spans:
            vals = _cells(c, row.get(c.key), tz)
            for k, v in enumerate(vals):
                cell = ws.cell(row=r, column=j + k, value=v)
                if v is not None:
                    if c.kind in ("pct", "pp"):
                        cell.number_format = PCT_FMT
                    elif c.kind == "datetime" and k == 0:
                        cell.number_format = "dd/mm/yyyy hh:mm"
                    elif c.kind == "min":
                        cell.number_format = "0.0"
            j += n
        motivo = "; ".join(f"{k}: {v}" for k, v in reasons.items() if v)
        notes = [NOTE_TEXT.get(n, n) for n in (row.get("_notes") or [])]
        cov = coverage_note(row.get("_coverage"))
        if cov:
            notes.append(cov)
        ws.cell(row=r, column=j, value=motivo or None)
        ws.cell(row=r, column=j + 1, value="; ".join(notes) or None)
        if cov:
            for jj in range(1, j + 2):
                ws.cell(row=r, column=jj).fill = WARN_FILL
        r += 1
    if t.rows:
        ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(headers))}{r - 1}"
    for j, h in enumerate(headers, start=1):
        letter = get_column_letter(j)
        ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, min(max(len(h) + 2, 12), 40))
    return r


def render(doc: Doc) -> bytes:
    tz = ZoneInfo(doc.tz)
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = doc.title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"{doc.scope_label} — {doc.period_text}"
    ws["A3"] = (f"gerado em {fmt.local_dt(doc.generated_at, tz)} · dados até {fmt.local_dt(doc.data_until, tz)} · "
                f"versão {doc.version}" + (f" — {doc.supersedes_text}" if doc.supersedes_text else ""))
    r = 4
    for b in doc.banners:
        ws.cell(row=r, column=1, value=b).fill = WARN_FILL
        r += 1
    r += 1
    if doc.summary_table:
        ws.cell(row=r, column=1, value=doc.summary_table.title).font = Font(bold=True)
        r = write_table(ws, doc.summary_table, tz, r + 1, human=True)
        if doc.summary_table.note:
            ws.cell(row=r, column=1, value=doc.summary_table.note)
            r += 1
    for o in doc.omitted:
        ws.cell(row=r, column=1, value=f"Seção omitida: {o}")
        r += 1
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 28)

    seen: set[int] = set()
    for t in doc.tables():
        if t is doc.summary_table or not t.sheet or id(t) in seen:
            continue
        seen.add(id(t))
        sh = wb.create_sheet(_sheet_name(wb, t.sheet))
        write_table(sh, t, tz, 1)
        sh.freeze_panes = "A2"

    rd = wb.create_sheet("Leia-me")
    rd["A1"] = "Definições e avisos"
    rd["A1"].font = Font(bold=True, size=12)
    rd.append([])
    rd.append(["Indicador", "Fórmula", "Referência"])
    for c in rd[3]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for d in doc.definitions:
        rd.append(list(d))
    rd.append([])
    for line in [*doc.banners, *[f"Seção omitida: {o}" for o in doc.omitted], *doc.readme,
                 "Valores brutos da API: frações com formato 0,0 %; tempos em minutos (coluna *_min); "
                 "datas em hora local da planta (*_local) e UTC ISO 8601 (*_utc).",
                 "Indicador indefinido = célula vazia + coluna 'motivo' (nunca 0).",
                 f"Versão do gerador: {doc.summary.get('_generator', '')}"]:
        if line:
            rd.append([line])
    rd.column_dimensions["A"].width = 32
    rd.column_dimensions["B"].width = 70
    rd.column_dimensions["C"].width = 32
    for row in rd.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
