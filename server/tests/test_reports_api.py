"""docs/reports.md — reports end to end on the real stack (Postgres + Influx + the API in process):

* every number of the report equals the API call that returns it (tolerance 1e-6), for daily_production, shift,
  weekly_oee and monthly_maintenance on the scenario C plant (real MQTT readings → detector → OEE);
* PDF/XLSX files, versions (§0.4), job mode, download;
* scheduling engine: never before send_at, never twice (two passes in the same minute), e-mail with attachments
  through an in-memory SMTP server, subject pattern, one-click unsubscribe, recipient capability filtering
  (operator never receives WO hours), no mail channel → admin warned, failure → retries 5/15/60 min → failed;
* tenant isolation of every endpoint.
"""
from __future__ import annotations

import asyncio
import io
import math
import uuid
from datetime import date, datetime, timedelta, timezone
from zoneinfo import ZoneInfo

import pytest
import pytest_asyncio
from openpyxl import load_workbook
from sqlalchemy import func, select, update

from conftest import _make_tenant
from oee_helpers import SCENARIO_C, publish_all, run_detector, scenario_c_readings, setup_scenario_c, ts
from report_helpers import MemorySMTP, attachments, body_text, find_link
from usage_helpers import api

TZ = ZoneInfo(SCENARIO_C["plant"]["timezone"])
DAY = ts(SCENARIO_C["shift"]["start_utc"]).astimezone(TZ).date()


def close(a, b, tol=1e-6):
    if a is None or b is None:
        return a is None and b is None
    if isinstance(a, bool) or isinstance(b, bool) or isinstance(a, str) or isinstance(b, str):
        return a == b
    return abs(a - b) <= tol * max(1.0, abs(b))


@pytest_asyncio.fixture(scope="session")
async def sc(app, client):
    """Scenario C plant (one line, the Torno CNC as output machine), stops justified by the operator."""
    from src.db import AsyncSessionLocal
    async with AsyncSessionLocal() as s:
        t = await _make_tenant(s, "RP")
        await s.commit()
    p, tags, shift, prod = await setup_scenario_c(client, t, label="RP")
    await publish_all(t.id, scenario_c_readings(tags))
    for minutes in (75, 150, 330, 481):
        await run_detector(p.machine["id"], t.id, ts(SCENARIO_C["shift"]["start_utc"]) + timedelta(minutes=minutes))
    start, end = ts(SCENARIO_C["shift"]["start_utc"]), ts(SCENARIO_C["shift"]["end_utc"])
    evs = await api(client, "GET", "/api/downtime/", t.headers("operator"),
                    params={"machine_id": p.machine["id"], "start": start.isoformat(), "stop": end.isoformat()})
    for e in evs:
        if e["source"] == "auto_run_state" and e["unjustified"] and e["duration_s"] and e["duration_s"] > 1200:
            await api(client, "PATCH", f"/api/downtime/{e['id']}", t.headers("operator"), json={"reason_code": "MAT-01"})
            break
    return {"t": t, "p": p, "shift": shift, "start": start, "end": end}


@pytest.fixture
def smtp(monkeypatch):
    from src.config import settings
    srv = MemorySMTP().start()
    monkeypatch.setattr(settings, "smtp_host", "127.0.0.1")
    monkeypatch.setattr(settings, "smtp_port", srv.port)
    monkeypatch.setattr(settings, "smtp_tls", False)
    monkeypatch.setattr(settings, "smtp_user", "")
    monkeypatch.setattr(settings, "smtp_from", "relatorios@factoryconnect.test")
    monkeypatch.setattr(settings, "sendgrid_api_key", "")
    yield srv
    srv.stop()


@pytest.fixture
def no_mail(monkeypatch):
    from src.config import settings
    monkeypatch.setattr(settings, "smtp_host", "")
    monkeypatch.setattr(settings, "sendgrid_api_key", "")


async def _gen(client, h, expect=200, **body):
    r = await client.post("/api/reports/generate", json=body, headers=h)
    assert r.status_code == expect, r.text
    return r


def _row(table, **match):
    for r in table["rows"]:
        if all(r.get(k) == v for k, v in match.items()):
            return r
    raise AssertionError(f"row {match} not in {table['title']}")


def _tables(doc):
    out = [doc["summary_table"]] if doc.get("summary_table") else []
    for s in doc["sections"]:
        out += [b for b in s["blocks"] if b["type"] == "table"]
    return out


# ── catalog / status ───────────────────────────────────────────────────────

async def test_catalog_and_status(client, tenants, no_mail):
    h = tenants.a.headers("operator")
    cat = await api(client, "GET", "/api/reports/catalog", h)
    assert [c["code"] for c in cat] == ["daily_production", "shift", "weekly_oee", "monthly_maintenance",
                                        "reliability_life"]
    by = {c["code"]: c for c in cat}
    assert by["daily_production"]["allowed"] and not by["monthly_maintenance"]["allowed"]   # operator
    assert by["shift"]["formats"] == ["pdf"] and by["weekly_oee"]["weekday"] == 0
    st = await api(client, "GET", "/api/reports/status", tenants.a.headers("admin"))
    assert st["mail_configured"] is False and "e-mail não configurado" in st["warning"]


# ── numbers = API (§0.1) ──────────────────────────────────────────────────

async def test_daily_production_numbers_equal_api(client, sc):
    t, p = sc["t"], sc["p"]
    h = t.headers("manager")
    doc = (await _gen(client, h, code="daily_production", scope_id=p.plant["id"], format="json",
                      params={"productive_day": DAY.isoformat()})).json()
    per = doc["summary"]["_period"]
    q = {"start": per["start"], "stop": per["end"], "group_by": "shift"}
    line = await api(client, "GET", "/api/oee", h, params={"scope": "line", "id": p.line["id"], **q})
    plant = await api(client, "GET", "/api/oee", h, params={"scope": "plant", "id": p.plant["id"], **q})
    assert line["oee"] is not None and 0 < line["oee"] < 1
    cover = doc["summary_table"]
    lrow = _row(cover, line_id=p.line["id"])
    prow = _row(cover, line_id=None)
    for k in ("oee", "availability", "performance", "quality", "produced", "good", "scrap",
              "cadence_planned_per_hour", "data_coverage"):
        assert close(lrow[k], line[k]), (k, lrow[k], line[k])
        assert close(prow[k], plant[k]), (k, prow[k], plant[k])
    # by shift: the points of the API, and the day total = the API total (S3-8)
    tables = _tables(doc)
    shift_tbl = next(x for x in tables if x["title"].startswith("OEE por turno (S3-8"))
    for pt in line["points"]:
        r = _row(shift_tbl, label=pt["label"])
        for k in ("oee", "availability", "performance", "quality", "good", "scrap", "pbt_s"):
            assert close(r[k], pt[k]), (k, r[k], pt[k])
    tot = _row(shift_tbl, label="Total do dia (soma dos turnos)")
    assert close(tot["oee"], line["oee"])
    # Pareto top-5 = the API Pareto
    par = await api(client, "GET", "/api/losses/pareto", h, params={"scope": "line", "id": p.line["id"],
                                                                     "start": per["start"], "stop": per["end"]})
    top = [r for r in par if r["duration_s"] > 0][:5]
    ptbl = next(x for x in tables if x["title"].startswith("Top-5"))
    assert [(r["label"], r["duration_s"], r["pct"]) for r in ptbl["rows"]] == \
           [(r["label"], r["duration_s"], r["pct"]) for r in top]
    # stops ≥ 10 min = the API downtime list of that productive day
    evs = await api(client, "GET", "/api/downtime/", h, params={"line_id": p.line["id"], "start": per["start"],
                                                                "stop": per["end"], "limit": 5000})
    long_ = sorted((e["start"], e["duration_s"]) for e in evs if e["productive_day"] == DAY.isoformat()
                   and not e["planned"] and (e["duration_s"] or 0) >= 600)
    stbl = next(x for x in tables if x["title"] == "Paradas ≥ 10 min")
    assert sorted((r["start"], r["duration_s"]) for r in stbl["rows"]) == long_ and long_
    # manager has maintenance.view → the maintenance section is there; every number came from a recorded API call
    assert any(s["title"] == "Manutenção do dia" for s in doc["sections"])
    paths = {c["path"] for c in doc["api_calls"]}
    assert {"/api/oee", "/api/losses/pareto", "/api/downtime/", "/api/failures/"} <= paths
    assert all(c["status"] == 200 for c in doc["api_calls"])

    # operator: same numbers, no maintenance section (no maintenance.view)
    op = (await _gen(client, t.headers("operator"), code="daily_production", scope_id=p.plant["id"], format="json",
                     params={"productive_day": DAY.isoformat()})).json()
    assert close(_row(op["summary_table"], line_id=p.line["id"])["oee"], line["oee"])
    assert not any(s["title"] == "Manutenção do dia" for s in op["sections"])
    assert any("maintenance.view" in o for o in op["omitted"])

    # XLSX: raw values of the API, fractions with 0.0 %, minutes, two timestamp columns, no formulas
    r = await _gen(client, h, code="daily_production", scope_id=p.plant["id"], format="xlsx",
                   params={"productive_day": DAY.isoformat()})
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["OEE_por_turno"]
    head = [c.value for c in ws[1]]
    row2 = [c.value for c in ws[2]]
    assert close(row2[head.index("oee")], line["points"][0]["oee"])
    assert ws.cell(row=2, column=head.index("oee") + 1).number_format == "0.0%"
    wsp = wb["Paradas"]
    ph = [c.value for c in wsp[1]]
    assert "inicio_local" in ph and "inicio_utc" in ph and "duracao_min" in ph
    for row in wsp.iter_rows(min_row=2, values_only=True):
        assert not any(isinstance(v, str) and v.startswith("=") for v in row)
    assert {"Resumo", "Leia-me", "Falhas", "OS"} <= set(wb.sheetnames)

    # PDF: versions — regenerating the same period makes version 2, the first stays downloadable
    r1 = await _gen(client, h, code="daily_production", scope_id=p.plant["id"], format="pdf",
                    params={"productive_day": DAY.isoformat()})
    r2 = await _gen(client, h, code="daily_production", scope_id=p.plant["id"], format="pdf",
                    params={"productive_day": DAY.isoformat()})
    assert r1.content[:5] == b"%PDF-" and len(r1.content) < 2 * 1024 * 1024
    v1, v2 = r1.headers["X-Report-Run-Id"], r2.headers["X-Report-Run-Id"]
    assert int(r2.headers["X-Report-Version"]) == int(r1.headers["X-Report-Version"]) + 1
    run2 = await api(client, "GET", f"/api/report-runs/{v2}", h)
    assert run2["supersedes_id"] == v1 and run2["period_key"] == DAY.isoformat()
    assert run2["data_until"] and run2["generator_version"] and len(run2["file_sha256"]) == 64
    d1 = await client.get(f"/api/report-runs/{v1}/download", headers=h)
    assert d1.status_code == 200 and d1.content == r1.content
    assert d1.headers["content-type"] == "application/pdf"
    # the operator's profile differs: they cannot download the manager's file
    assert (await client.get(f"/api/report-runs/{v1}/download", headers=t.headers("operator"))).status_code == 404


async def test_shift_weekly_monthly_numbers_equal_api(client, sc):
    t, p = sc["t"], sc["p"]
    h = t.headers("manager")
    sh = (await _gen(client, h, code="shift", scope_type="line", scope_id=p.line["id"], format="json",
                     params={"productive_day": DAY.isoformat(), "shift_id": sc["shift"]["id"]})).json()
    sr = await api(client, "GET", "/api/reports/shift", h, params={"line_id": p.line["id"],
                                                                     "productive_day": DAY.isoformat(),
                                                                     "shift_id": sc["shift"]["id"]})
    row = sh["summary_table"]["rows"][0]
    for k in ("oee", "availability", "performance", "quality", "good", "scrap", "data_coverage"):
        assert close(row[k], sr["oee"][k]), k
    prod = next(x for x in _tables(sh) if x["title"] == "Produção por máquina")
    assert [(r["machine_id"], r["good"], r["oee"]) for r in prod["rows"]] == \
           [(m["machine_id"], m["good"], m["oee"]) for m in sr["production"]]
    assert any(s["title"] == "Pendências para o próximo turno" for s in sh["sections"])
    assert (await _gen(client, h, code="shift", scope_type="line", scope_id=p.line["id"], format="xlsx",
                       expect=422, params={"productive_day": DAY.isoformat()}))

    wk = (await _gen(client, h, code="weekly_oee", scope_type="line", scope_id=p.line["id"], format="json",
                     params={"week_start": DAY.isoformat()})).json()
    wper = wk["summary"]["_period"]
    wo = await api(client, "GET", "/api/oee", h, params={"scope": "line", "id": p.line["id"], "start": wper["start"],
                                                          "stop": wper["end"], "group_by": "shift"})
    assert close(wk["summary_table"]["rows"][0]["oee"], wo["oee"]) and wo["oee"] is not None
    assert len(wk["summary"]["trend"]) == 12
    sn = await api(client, "GET", "/api/oee", h, params={"scope": "line", "id": p.line["id"], "start": wper["start"],
                                                          "stop": wper["end"], "group_by": "shift_name"})
    tbl = next(x for x in _tables(wk) if x["title"].startswith("OEE por turno"))
    assert [(r["label"], r["oee"]) for r in tbl["rows"]] == [(x["label"], x["oee"]) for x in sn["points"]]
    pdf = await _gen(client, h, code="weekly_oee", scope_type="line", scope_id=p.line["id"], format="pdf",
                     params={"week_start": DAY.isoformat()})
    assert pdf.content[:5] == b"%PDF-"

    month = f"{DAY.year}-{DAY.month:02d}"
    mm = (await _gen(client, h, code="monthly_maintenance", scope_id=p.plant["id"], format="json",
                     params={"month": month})).json()
    first = date(DAY.year, DAY.month, 1)
    nxt = date(DAY.year + (DAY.month == 12), DAY.month % 12 + 1, 1)
    s = datetime.combine(first, datetime.min.time(), TZ).astimezone(timezone.utc)
    e = datetime.combine(nxt, datetime.min.time(), TZ).astimezone(timezone.utc)
    kp = await api(client, "GET", "/api/maintenance/kpis", h, params={"plant_id": p.plant["id"],
                                                                        "start": s.isoformat(), "stop": e.isoformat()})
    cur = mm["summary_table"]["rows"][0]
    for k in ("pmp", "schedule_compliance", "backlog_weeks"):
        assert close(cur[k], kp[k]), k
    rel = await api(client, "GET", "/api/reliability", h, params={
        "scope": "plant", "id": p.plant["id"], "stop": e.isoformat(),
        "start": datetime.combine(_months_before(first, 11), datetime.min.time(), TZ).astimezone(timezone.utc).isoformat()})
    assert close(cur["mtbf_h"], rel["mtbf_h"]) and cur["failures"] == rel["failures"]
    assert rel["failures"] >= 1                                # scenario C faults became failures
    x = await _gen(client, h, code="monthly_maintenance", scope_id=p.plant["id"], format="xlsx", params={"month": month})
    assert "OS" in load_workbook(io.BytesIO(x.content)).sheetnames
    # operator: no maintenance.kpis → 403 for the monthly report (screen-equivalent permission)
    await _gen(client, t.headers("operator"), expect=403, code="monthly_maintenance", scope_id=p.plant["id"],
               format="json", params={"month": month})


def _months_before(d: date, k: int) -> date:
    idx = d.year * 12 + d.month - 1 - k
    return date(idx // 12, idx % 12 + 1, 1)


async def test_job_mode_and_undefined_values(client, sc):
    t, p = sc["t"], sc["p"]
    h = t.headers("manager")
    r = await _gen(client, h, expect=202, code="daily_production", scope_id=p.plant["id"], format="pdf",
                   mode="job", params={"productive_day": DAY.isoformat()})
    rid = r.json()["run_id"]
    for _ in range(200):
        run = await api(client, "GET", f"/api/report-runs/{rid}", h)
        if run["status"] != "generating":
            break
        await asyncio.sleep(0.1)
    assert run["status"] == "done", run
    d = await client.get(f"/api/report-runs/{rid}/download", headers=h)
    assert d.status_code == 200 and d.content[:5] == b"%PDF-"
    # a day without production: OEE undefined → null + reason (never 0)
    empty_day = (DAY - timedelta(days=20)).isoformat()
    doc = (await _gen(client, h, code="daily_production", scope_id=p.plant["id"], format="json",
                      params={"productive_day": empty_day})).json()
    lrow = _row(doc["summary_table"], line_id=p.line["id"])
    assert lrow["oee"] is None and lrow["_reason"]["oee"]
    x = await _gen(client, h, code="daily_production", scope_id=p.plant["id"], format="xlsx",
                   params={"productive_day": empty_day})
    ws = load_workbook(io.BytesIO(x.content))["Resumo"]
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    hdr = next(i for i, r in enumerate(rows) if r and r[0] == "Linha")
    oee_i, mot_i = rows[hdr].index("OEE"), rows[hdr].index("motivo")
    line_row = next(r for r in rows[hdr + 1:] if r[0] == p.line["name"])
    assert line_row[oee_i] is None and line_row[mot_i] and "oee" in line_row[mot_i]


# ── Weibull API (§2) ──────────────────────────────────────────────────────

async def _weibull_plant(client, t, cutoff: datetime):
    """weibull_cylinders: 6 failures, 2 preventive removals, 2 in service (C-09 2.0 M, C-10 5.0 M at 20 000/day)."""
    from src.db import AsyncSessionLocal
    from src.models import Component, ComponentType, Position, RemovalReason, UsageBucket, UsageTotal
    from usage_helpers import build_plant
    p = await build_plant(client, t, f"W{uuid.uuid4().hex[:4]}")
    rows = [("C-01", 3100000, "failure"), ("C-02", 3900000, "failure"), ("C-03", 4400000, "failure"),
            ("C-04", 5200000, "failure"), ("C-05", 5800000, "failure"), ("C-06", 6700000, "failure"),
            ("C-07", 4000000, "preventive_replacement"), ("C-08", 4500000, "preventive_replacement"),
            ("C-09", 2000000, None), ("C-10", 5000000, None)]
    async with AsyncSessionLocal() as s:
        for i, (sn, age, why) in enumerate(rows):
            pos = Position(tenant_id=t.id, machine_id=p.machine["id"], name=f"Cilindro {i}",
                           component_type=ComponentType.pneumatic_cylinder)
            s.add(pos)
            await s.flush()
            installed = cutoff - timedelta(days=400)
            c = Component(tenant_id=t.id, component_type=ComponentType.pneumatic_cylinder, manufacturer="Festo",
                          model="DSBC-32", serial_number=sn, nominal_life={"cycles": 5000000},
                          installed_at=installed, last_position_id=pos.id)
            if why:
                c.removed_at = cutoff - timedelta(days=30)
                c.removal_reason = RemovalReason(why)
            else:
                c.position_id = pos.id
            s.add(c)
            await s.flush()
            if why:
                s.add(UsageTotal(scope="component", entity_id=c.id, measure="cycles", tenant_id=t.id, value=age))
            else:
                h0 = installed.replace(minute=0, second=0, microsecond=0)
                s.add(UsageBucket(scope="component", entity_id=c.id, measure="cycles", tenant_id=t.id,
                                  bucket_start=h0 + timedelta(hours=1), value=age - 14 * 20000))
                for d in range(14):
                    b = (cutoff - timedelta(days=d + 1)).replace(minute=0, second=0, microsecond=0)
                    s.add(UsageBucket(scope="component", entity_id=c.id, measure="cycles", tenant_id=t.id,
                                      bucket_start=b, value=20000))
                s.add(UsageTotal(scope="component", entity_id=c.id, measure="cycles", tenant_id=t.id, value=age))
        await s.commit()
    return p


async def test_reliability_life_weibull_api_and_report(client):
    import json
    import re
    from pathlib import Path
    from src.db import AsyncSessionLocal
    doc_md = (Path(__file__).resolve().parents[2] / "docs" / "reports.md").read_text(encoding="utf-8")
    cases = {c["id"]: c for c in json.loads(re.search(r"```json report-weibull-cases\n(.*?)\n```", doc_md, re.S).group(1))}
    exp = cases["weibull_cylinders"]["expected"]
    async with AsyncSessionLocal() as s:
        ta = await _make_tenant(s, "WA")
        tb = await _make_tenant(s, "WB")
        await s.commit()
    cutoff = datetime.now(timezone.utc).replace(microsecond=0) - timedelta(hours=2)
    p = await _weibull_plant(client, ta, cutoff)
    h = ta.headers("manager")
    data = await api(client, "GET", "/api/reports/reliability-life", h,
                     params={"plant_id": p.plant["id"], "cutoff": cutoff.isoformat()})
    [typ] = data["types"]
    assert typ["measure"] == "cycles"
    [g] = typ["groups"]
    assert g["n"] == exp["n"] and g["r"] == exp["r"] and g["suspensions_in_service"] == 2
    for k in ("beta", "eta", "se_beta", "se_eta", "beta_lo", "beta_hi", "eta_lo", "eta_hi", "mttf", "b10", "b10_lo",
              "b10_hi", "b10_over_nominal", "simple_mttf"):
        assert math.isclose(g[k], exp[k], rel_tol=1e-6), (k, g[k], exp[k])
    assert g["pattern"] == exp["pattern"] and g["low_confidence"] is True and g["below_nominal"] is True
    risk = {r["age"]: r["p_fail"] for r in data["in_service_risk"]}
    for cond in cases["weibull_cylinders"]["conditional"]:
        assert math.isclose(risk[cond["age"]], cond["p_fail"], abs_tol=5e-7), (cond, risk)
    assert data["in_service_risk"][0]["age"] == 5000000          # ordered by P desc
    # the report = that API (numbers copied, never recomputed)
    rep = (await _gen(client, h, code="reliability_life", scope_id=p.plant["id"], format="json",
                      params={"cutoff_date": None})).json()
    grow = rep["summary_table"]["rows"][0]
    assert {k: grow[k] for k in ("beta", "eta", "b10", "mttf")} == {k: g[k] for k in ("beta", "eta", "b10", "mttf")} or \
        all(math.isclose(grow[k], g[k], rel_tol=1e-9) for k in ("beta", "eta", "b10", "mttf"))
    x = await _gen(client, h, code="reliability_life", scope_id=p.plant["id"], format="xlsx")
    wb = load_workbook(io.BytesIO(x.content))
    assert {"Weibull_grupos", "Weibull_unidades", "Leia-me", "Resumo"} <= set(wb.sheetnames)
    wu = wb["Weibull_unidades"]
    hdr = [c.value for c in wu[1]]
    units = [dict(zip(hdr, [c.value for c in r])) for r in wu.iter_rows(min_row=2)]
    assert len(units) == 10 and sum(u["delta"] for u in units) == 6
    assert sorted(u["idade"] for u in units) == sorted([3100000, 3900000, 4400000, 5200000, 5800000, 6700000,
                                                         4000000, 4500000, 2000000, 5000000])
    pdf = await _gen(client, h, code="reliability_life", scope_id=p.plant["id"], format="pdf")
    assert pdf.content[:5] == b"%PDF-"
    # isolation: tenant B never sees A's components, and A's plant is "not found" for B
    hb = tb.headers("manager")
    assert (await client.get("/api/reports/reliability-life", headers=hb,
                             params={"plant_id": p.plant["id"]})).status_code == 404
    other = await api(client, "GET", "/api/reports/reliability-life", hb)
    assert other["types"] == [] and other["in_service_risk"] == []
    # operator: no maintenance reliability screen → 403
    assert (await client.get("/api/reports/reliability-life", headers=ta.headers("operator"))).status_code == 403


# ── tenant isolation ──────────────────────────────────────────────────────

async def test_tenant_isolation(client, sc, tenants):
    t, p = sc["t"], sc["p"]
    ha, hb = t.headers("manager"), tenants.b.headers("manager")
    s = await api(client, "POST", "/api/report-schedules/", ha, 201, json={
        "report_type": "daily_production", "scope_type": "plant", "scope_id": p.plant["id"]})
    run = await _gen(client, ha, code="daily_production", scope_id=p.plant["id"], format="pdf",
                     params={"productive_day": DAY.isoformat()})
    rid = run.headers["X-Report-Run-Id"]
    sid = s["id"]
    for method, url in (("GET", f"/api/report-schedules/{sid}"), ("PATCH", f"/api/report-schedules/{sid}"),
                        ("DELETE", f"/api/report-schedules/{sid}"), ("POST", f"/api/report-schedules/{sid}/pause"),
                        ("POST", f"/api/report-schedules/{sid}/resume"),
                        ("POST", f"/api/report-schedules/{sid}/send-now"),
                        ("GET", f"/api/report-schedules/{sid}/history"), ("GET", f"/api/report-runs/{rid}"),
                        ("GET", f"/api/report-runs/{rid}/download")):
        kw = {"json": {"name": "x"}} if method == "PATCH" else {}
        r = await client.request(method, url, headers=hb, **kw)
        assert r.status_code == 404, (method, url, r.status_code)
    assert sid not in {x["id"] for x in await api(client, "GET", "/api/report-schedules/", hb)}
    assert rid not in {x["id"] for x in await api(client, "GET", "/api/report-runs/", hb)}
    # B cannot point at A's plant, line or users
    r = await client.post("/api/report-schedules/", headers=hb, json={"report_type": "daily_production",
                                                                      "scope_type": "plant", "scope_id": p.plant["id"]})
    assert r.status_code == 404
    r = await client.post("/api/report-schedules/", headers=tenants.b.headers("admin"), json={
        "report_type": "reliability_life", "scope_type": "tenant",
        "recipients": {"roles": ["manager"], "user_ids": [t.users["manager"]["id"]]}})
    assert r.status_code == 404
    await _gen(client, hb, expect=404, code="daily_production", scope_id=p.plant["id"], format="json")
    await _gen(client, hb, expect=404, code="shift", scope_type="line", scope_id=p.line["id"], format="json")
    r = await client.put("/api/reports/line-targets", headers=hb, json={"line_id": p.line["id"], "oee": 0.8,
                                                                        "valid_from": DAY.isoformat()})
    assert r.status_code == 404
    # roles: operator cannot manage schedules
    assert (await client.get("/api/report-schedules/", headers=t.headers("operator"))).status_code == 403
    await api(client, "DELETE", f"/api/report-schedules/{sid}", ha, 204)


# ── scheduling engine + e-mail ─────────────────────────────────────────────

async def _backdate(schedule_id: str, days: int = 60, to=None):
    from src.db import AsyncSessionLocal
    from src.models import ReportSchedule
    async with AsyncSessionLocal() as s:
        await s.execute(update(ReportSchedule).where(ReportSchedule.id == schedule_id).values(
            created_at=to if to is not None else func.now() - timedelta(days=days)))
        await s.commit()


async def _tick(now, tenant_id, app):
    from src.db import AsyncSessionLocal
    from src.services.reports import engine
    return await engine.tick(AsyncSessionLocal, now=now, app=app, tenant_id=tenant_id)


async def _count(model, **where):
    from src.db import AsyncSessionLocal
    async with AsyncSessionLocal() as s:
        q = select(func.count()).select_from(model)
        for k, v in where.items():
            q = q.where(getattr(model, k) == v)
        return (await s.execute(q)).scalar_one()


async def test_schedule_engine_sends_once_with_attachments_and_unsubscribe(client, app, sc, smtp):
    from src.models import ReportDelivery, ReportDispatch
    from src.services.oee.engine import load_calendar
    from src.services.reports import periods as P
    from src.db import AsyncSessionLocal
    t, p = sc["t"], sc["p"]
    s = await api(client, "POST", "/api/report-schedules/", t.headers("manager"), 201, json={
        "report_type": "daily_production", "scope_type": "plant", "scope_id": p.plant["id"],
        "formats": ["pdf", "xlsx"], "deliver_at_local": "06:30"})
    assert s["recipients"]["roles"] == ["manager", "production_supervisor", "admin"]     # §5.1 default
    assert s["mail_configured"] is True and s["warning"] is None and s["next_send_at"]
    async with AsyncSessionLocal() as db:
        cal = await load_calendar(db, t.id, p.plant["id"], TZ)
    per = P.day_period(cal, DAY)
    send = P.send_at(per, "daily_production", TZ, "06:30")
    # created just before that send: earlier periods are not back-filled (only DAY is due)
    await _backdate(s["id"], to=send - timedelta(hours=1))
    ready = P.ready_at(per, "daily_production")
    assert send == max(ready, datetime.combine(DAY + timedelta(days=1), datetime.min.time().replace(hour=6, minute=30),
                                               TZ).astimezone(timezone.utc))
    # 1 minute before send_at: nothing for that period
    await _tick(send - timedelta(minutes=1), t.id, app)
    assert await _count(ReportDispatch, schedule_id=s["id"], period_key=DAY.isoformat()) == 0
    # at send_at, twice in the same minute (two engine passes) → one dispatch, one e-mail per recipient
    await asyncio.gather(_tick(send, t.id, app), _tick(send, t.id, app))
    await _tick(send + timedelta(seconds=30), t.id, app)
    assert await _count(ReportDispatch, schedule_id=s["id"], period_key=DAY.isoformat()) == 1
    mgr, adm = t.users["manager"]["email"], t.users["admin"]["email"]
    assert len(smtp.to(mgr)) == 1 and len(smtp.to(adm)) == 1
    assert smtp.to(t.users["operator"]["email"]) == []           # operator: not a default recipient
    msg = smtp.to(mgr)[0]
    assert msg["Subject"] == f"[FactoryConnect] Produção diária — {p.plant['name']} — {DAY:%d/%m/%Y}"
    att = attachments(msg)
    assert any(n.endswith(".pdf") and c[:5] == b"%PDF-" for n, c in att.items())
    assert any(n.endswith(".xlsx") for n in att)
    text = body_text(msg)
    assert "OEE da planta" in text and "Abrir no FactoryConnect" in text
    assert msg["List-Unsubscribe"]
    unsub = find_link(text, "/api/reports/unsubscribe?token=")
    hist = await api(client, "GET", f"/api/report-schedules/{s['id']}/history", t.headers("manager"))
    d = next(x for x in hist["dispatches"] if x["period_key"] == DAY.isoformat())
    assert d["status"] == "sent" and d["send_at"] == send.isoformat().replace("+00:00", "Z")
    assert {x["email"] for x in hist["deliveries"] if x["status"] == "sent"} >= {mgr, adm}
    assert all(x["provider_id"] for x in hist["deliveries"] if x["status"] == "sent")
    assert {r["format"] for r in hist["runs"]} == {"pdf", "xlsx"}
    # one-click unsubscribe (no login) → the next send skips that address
    path = unsub.split("factoryconnect.tech", 1)[-1] if "://" in unsub else unsub
    path = "/api" + path.split("/api", 1)[1]
    r = await client.get(path)
    assert r.status_code == 200 and "Descadastro confirmado" in r.text
    smtp.clear()
    res = await api(client, "POST", f"/api/report-schedules/{s['id']}/send-now", t.headers("manager"))
    assert res["status"] == "sent"
    assert smtp.to(mgr) == [] and len(smtp.to(adm)) == 1
    assert await _count(ReportDelivery, dispatch_id=res["dispatch_id"], email=mgr, status="unsubscribed") == 1
    # pause → no new dispatch for the next day
    await api(client, "POST", f"/api/report-schedules/{s['id']}/pause", t.headers("manager"))
    nxt = P.day_period(cal, DAY + timedelta(days=1))
    await _tick(P.send_at(nxt, "daily_production", TZ, "06:30") + timedelta(minutes=1), t.id, app)
    assert await _count(ReportDispatch, schedule_id=s["id"], period_key=nxt.key) == 0
    await api(client, "DELETE", f"/api/report-schedules/{s['id']}", t.headers("manager"), 204)


async def test_operator_never_receives_wo_hours(client, app, sc, smtp):
    t, p = sc["t"], sc["p"]
    s = await api(client, "POST", "/api/report-schedules/", t.headers("admin"), 201, json={
        "report_type": "monthly_maintenance", "scope_type": "plant", "scope_id": p.plant["id"],
        "formats": ["pdf", "xlsx"],
        "recipients": {"roles": ["manager"], "user_ids": [t.users["operator"]["id"]]}})
    res = await api(client, "POST", f"/api/report-schedules/{s['id']}/send-now", t.headers("admin"))
    assert res["status"] == "sent"
    op_msgs = smtp.to(t.users["operator"]["email"])
    mg_msgs = smtp.to(t.users["manager"]["email"])
    assert len(op_msgs) == 1 and len(mg_msgs) == 1
    op_x = next(c for n, c in attachments(op_msgs[0]).items() if n.endswith(".xlsx"))
    mg_x = next(c for n, c in attachments(mg_msgs[0]).items() if n.endswith(".xlsx"))
    op_wb, mg_wb = load_workbook(io.BytesIO(op_x)), load_workbook(io.BytesIO(mg_x))
    assert "OS" in mg_wb.sheetnames and "Indicadores" in mg_wb.sheetnames
    assert "OS" not in op_wb.sheetnames and "Indicadores" not in op_wb.sheetnames
    assert "maintenance.kpis" in body_text(op_msgs[0])
    hist = await api(client, "GET", f"/api/report-schedules/{s['id']}/history", t.headers("admin"))
    assert len({r["id"] for r in hist["runs"]}) == 4                  # 2 profiles × 2 formats
    await api(client, "DELETE", f"/api/report-schedules/{s['id']}", t.headers("admin"), 204)


async def test_no_mail_channel_warns_admin_and_keeps_download(client, app, sc, no_mail):
    t, p = sc["t"], sc["p"]
    s = await api(client, "POST", "/api/report-schedules/", t.headers("manager"), 201, json={
        "report_type": "shift", "scope_type": "line", "scope_id": p.line["id"],
        "params": {"shift_id": sc["shift"]["id"]}, "recipients": {"roles": ["production_supervisor", "manager"]}})
    assert s["mail_configured"] is False and "e-mail não configurado" in s["warning"]
    assert s["deliver_at_local"] is None
    res = await api(client, "POST", f"/api/report-schedules/{s['id']}/send-now", t.headers("manager"))
    assert res["status"] == "no_channel"
    hist = await api(client, "GET", f"/api/report-schedules/{s['id']}/history", t.headers("manager"))
    assert hist["runs"] and all(r["status"] == "done" for r in hist["runs"])
    assert {x["status"] for x in hist["deliveries"]} == {"no_channel"}
    notes = await api(client, "GET", "/api/notifications/", t.headers("admin"))
    items = notes["items"] if isinstance(notes, dict) else notes
    assert any("E-mail não configurado" in n["title"] for n in items)
    rid = hist["runs"][0]["id"]
    assert (await client.get(f"/api/report-runs/{rid}/download", headers=t.headers("manager"))).status_code == 200
    await api(client, "DELETE", f"/api/report-schedules/{s['id']}", t.headers("manager"), 204)


async def test_generation_failure_retries_then_failed(client, app, sc, smtp, monkeypatch):
    from src.models import ReportDispatch, ReportRun
    from src.services.reports import builders
    from src.db import AsyncSessionLocal
    t, p = sc["t"], sc["p"]
    s = await api(client, "POST", "/api/report-schedules/", t.headers("manager"), 201, json={
        "report_type": "weekly_oee", "scope_type": "plant", "scope_id": p.plant["id"]})
    await _backdate(s["id"])

    async def boom(src, ctx):
        raise RuntimeError("influx down")
    monkeypatch.setitem(builders.BUILDERS, "weekly_oee", boom)
    now = datetime.now(timezone.utc)
    await _tick(now, t.id, app)
    async with AsyncSessionLocal() as db:
        ds = (await db.execute(select(ReportDispatch).where(ReportDispatch.schedule_id == s["id"]))).scalars().all()
    assert ds, "a due weekly period"
    d = max(ds, key=lambda x: x.send_at)
    waits = []
    cur = now
    for _ in range(3):
        async with AsyncSessionLocal() as db:
            row = await db.get(ReportDispatch, d.id)
        assert row.status == "pending" and "influx down" in row.last_error
        waits.append(round((row.next_attempt_at - cur).total_seconds() / 60))
        await _tick(row.next_attempt_at - timedelta(seconds=30), t.id, app)    # not yet
        async with AsyncSessionLocal() as db:
            assert (await db.get(ReportDispatch, d.id)).attempts == row.attempts
        cur = row.next_attempt_at
        await _tick(cur, t.id, app)
    assert waits == [5, 15, 60]
    async with AsyncSessionLocal() as db:
        row = await db.get(ReportDispatch, d.id)
        failed = (await db.execute(select(ReportRun).where(ReportRun.dispatch_id == d.id,
                                                           ReportRun.status == "failed"))).scalars().all()
    assert row.status == "failed" and row.attempts == 4 and failed
    assert smtp.messages == []                                       # nothing sent partially
    notes = await api(client, "GET", "/api/notifications/", t.headers("admin"))
    items = notes["items"] if isinstance(notes, dict) else notes
    assert any("Falha ao gerar o relatório" in n["title"] for n in items)
    await api(client, "DELETE", f"/api/report-schedules/{s['id']}", t.headers("manager"), 204)


async def test_large_attachment_link_only_and_external_bounces(client, app, sc, smtp, monkeypatch):
    from src.config import settings
    t, p = sc["t"], sc["p"]
    ha = t.headers("admin")
    # externals need the admin switch (LGPD)
    r = await client.post("/api/reports/external-recipients", headers=ha, json={
        "email": "auditor@cliente-exemplo.com.br", "name": "Auditor", "purpose": "auditoria mensal", "legal_basis": "contract"})
    assert r.status_code == 409
    await api(client, "PUT", "/api/reports/settings", ha, json={"external_recipients_enabled": True})
    ext = await api(client, "POST", "/api/reports/external-recipients", ha, 201, json={
        "email": "auditor@cliente-exemplo.com.br", "name": "Auditor", "purpose": "auditoria mensal", "legal_basis": "contract"})
    s = await api(client, "POST", "/api/report-schedules/", ha, 201, json={
        "report_type": "shift", "scope_type": "line", "scope_id": p.line["id"], "params": {"shift_id": sc["shift"]["id"]},
        "recipients": {"roles": ["manager"], "external_ids": [ext["id"]]}})
    monkeypatch.setattr(settings, "reports_attachment_max_bytes", 100)
    res = await api(client, "POST", f"/api/report-schedules/{s['id']}/send-now", ha)
    assert res["status"] == "sent"
    msg = smtp.to(t.users["manager"]["email"])[0]
    assert attachments(msg) == {}
    link = find_link(body_text(msg), "/api/report-runs/")
    assert link and "link=" in link
    path = "/api" + link.split("/api", 1)[1]
    d = await client.get(path, headers=t.headers("manager"))
    assert d.status_code == 200 and d.content[:5] == b"%PDF-"
    assert (await client.get(path.split("link=")[0] + "link=forged.token.x", headers=t.headers("manager"))).status_code == 410
    assert len(smtp.to("auditor@cliente-exemplo.com.br")) == 1
    # 3 bounces disable the external recipient and warn the admin
    smtp.refuse.add("auditor@cliente-exemplo.com.br")
    for _ in range(3):
        await api(client, "POST", f"/api/report-schedules/{s['id']}/send-now", ha)
    exts = await api(client, "GET", "/api/reports/external-recipients", ha)
    e = next(x for x in exts if x["id"] == ext["id"])
    assert e["bounce_count"] == 3 and e["disabled_at"]
    hist = await api(client, "GET", f"/api/report-schedules/{s['id']}/history", ha)
    assert sum(1 for x in hist["deliveries"] if x["status"] == "bounced") == 3
    await api(client, "DELETE", f"/api/report-schedules/{s['id']}", ha, 204)
    await api(client, "PUT", "/api/reports/settings", ha, json={"external_recipients_enabled": False})


async def test_line_target_column_appears(client, sc):
    t, p = sc["t"], sc["p"]
    await api(client, "PUT", "/api/reports/line-targets", t.headers("manager"), json={
        "line_id": p.line["id"], "oee": 0.85, "valid_from": (DAY - timedelta(days=1)).isoformat()})
    doc = (await _gen(client, t.headers("manager"), code="daily_production", scope_id=p.plant["id"], format="json",
                      params={"productive_day": DAY.isoformat()})).json()
    assert "target_oee" in {c["key"] for c in doc["summary_table"]["columns"]}
    assert _row(doc["summary_table"], line_id=p.line["id"])["target_oee"] == 0.85
    from src.db import AsyncSessionLocal
    from src.models import LineTarget
    async with AsyncSessionLocal() as s:
        await s.execute(LineTarget.__table__.delete().where(LineTarget.line_id == p.line["id"]))
        await s.commit()
